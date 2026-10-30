"""
Comprehensive Source Code Compendium Compiler
Generates both 'Source_Code_Compendium.docx' and 'Source_Code_Compendium.pdf'
Contains 100% complete unabridged code for all project files.
"""

import os
import sys
from datetime import datetime

# -------------------------------------------------------------
# Document Specification & File Manifest
# -------------------------------------------------------------
PROJECT_TITLE = "Credit Card Fraud Detection Using Anomaly Detection"
PROJECT_NO = "Project No. 32"
PROGRAMME = "HCLTech Career Shaper Programme — III Year Industry Project"
LEARNER = "Saksham"
COORDINATOR = "Shaily Gaur"
GITHUB_REPO = "https://github.com/siwalsakshi/Credit-Card-Fraud-Detection-Using-Anomaly-Detection"
LIVE_VERCEL = "https://credit-card-fraud-detection-rho-ten.vercel.app?_vercel_share=5uKjBkdTqAg2QJC7G7Apv8vuOwbrPPA5"

FILES_MANIFEST = [
    {
        "section": "1. Full-Stack Web Application Server & REST API",
        "path": "app.py",
        "description": "Production Python HTTP server providing REST API inference for 4 anomaly models, dataset querying, fairness audits, and static file serving.",
        "lang": "Python"
    },
    {
        "section": "2. Machine Learning Anomaly Detection Pipeline & Fairness Audit",
        "path": "anomaly_detection.py",
        "description": "End-to-end ML pipeline training Isolation Forest, Local Outlier Factor, Neural Autoencoder, and One-Class SVM with SMOTE and ethical FPR audit.",
        "lang": "Python"
    },
    {
        "section": "3. Synthetic Dataset Generator & SMOTE Augmentation",
        "path": "generate_dataset.py",
        "description": "Generates 5,000 calibrated payment transactions with cyclical diurnal features, micro-testing probes, and log amount distributions.",
        "lang": "Python"
    },
    {
        "section": "4. Exploratory Data Analysis & Visualisation Generator",
        "path": "eda_analysis.py",
        "description": "Generates publication-quality charts for class imbalance, log-amount distributions, temporal heatmaps, and category breakdowns.",
        "lang": "Python"
    },
    {
        "section": "5. System Architecture & DFD Diagram Generator",
        "path": "generate_architecture_diagram.py",
        "description": "Generates Level 0 Data Flow Diagram (DFD) and 4-tier system architecture diagrams using Matplotlib.",
        "lang": "Python"
    },
    {
        "section": "6. Vercel Serverless Function Engine",
        "path": "api/index.py",
        "description": "Vercel Serverless Function entrypoint wrapping BaseHTTPRequestHandler with warm model memory caching and multi-route dispatching.",
        "lang": "Python"
    },
    {
        "section": "7. Frontend User Interface — Glassmorphism Dashboard",
        "path": "web/index.html",
        "description": "Interactive web dashboard featuring real-time risk gauges, preset attack scenarios, model comparison grid, and transaction stream.",
        "lang": "HTML"
    },
    {
        "section": "8. Frontend Design System & Responsive Stylesheet",
        "path": "web/styles.css",
        "description": "Dark-mode glassmorphism design system featuring CSS variables, animated SVG gauges, responsive grids, and typography.",
        "lang": "CSS"
    },
    {
        "section": "9. Frontend Client Logic & Reactive Gauges",
        "path": "web/app.js",
        "description": "Client-side controller managing async REST API calls, dynamic radial gauge animations, preset loading, and pagination.",
        "lang": "JavaScript"
    },
    {
        "section": "10. Data Dictionary & Transaction Schema",
        "path": "data_dictionary.csv",
        "description": "Column definitions, data types, nullability constraints, domain ranges, and business descriptions for payment records.",
        "lang": "CSV"
    },
    {
        "section": "11. Cloud & Container Deployment Configurations",
        "path": "Dockerfile",
        "description": "Multi-layer production container definition for automated Docker builds.",
        "lang": "Dockerfile"
    },
    {
        "section": "12. Vercel Serverless Routing Specification",
        "path": "vercel.json",
        "description": "Serverless function inclusion patterns and path rewrites for Vercel deployment.",
        "lang": "JSON"
    },
    {
        "section": "13. Cloud Blueprint & Process Configuration",
        "path": "render.yaml",
        "description": "Render 1-click infrastructure blueprint specification for web services.",
        "lang": "YAML"
    },
    {
        "section": "14. Process Control & Dependency Files",
        "path": "Procfile",
        "description": "Universal process launcher for web servers (Render, Railway, Heroku).",
        "lang": "Text"
    },
    {
        "section": "15. Core Project Dependencies",
        "path": "requirements.txt",
        "description": "Python package dependency manifest with pinned version minimums.",
        "lang": "Text"
    }
]

# -------------------------------------------------------------
# Part 1: DOCX Generation via python-docx
# -------------------------------------------------------------
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_docx():
    print("[DOCX] Initializing document...")
    doc = docx.Document()
    
    # Page setup: Standard Letter, 0.75-inch margins
    for sec in doc.sections:
        sec.top_margin = Inches(0.75)
        sec.bottom_margin = Inches(0.75)
        sec.left_margin = Inches(0.75)
        sec.right_margin = Inches(0.75)

    # 1. Title Block
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(20)
    p_title.paragraph_format.space_after = Pt(4)
    run_t = p_title.add_run("PROJECT SOURCE CODE COMPENDIUM")
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(22)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(16)
    run_sub = p_sub.add_run(f"{PROJECT_TITLE} ({PROJECT_NO})\n{PROGRAMME}")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Meta Table
    tbl_meta = doc.add_table(rows=4, cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_rows = [
        ("Student / Learner", f"{LEARNER} (CSE / AIML)"),
        ("Project Coordinator", COORDINATOR),
        ("Official GitHub Repo", GITHUB_REPO),
        ("Live Vercel Deployment", LIVE_VERCEL)
    ]
    for i, (k, v) in enumerate(meta_rows):
        row = tbl_meta.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.8)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(30, 41, 59)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        if "http" in v:
            r1.font.color.rgb = RGBColor(2, 132, 199)
            r1.font.underline = True
        else:
            r1.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # System Architecture Diagram
    if os.path.exists("visualizations/00_system_architecture.png"):
        h_arch = doc.add_heading("System Architecture & Data Flow Diagram (Level 0)", level=1)
        h_arch.paragraph_format.space_after = Pt(8)
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_after = Pt(16)
        p_img.add_run().add_picture("visualizations/00_system_architecture.png", width=Inches(6.5))

    # Table of Contents Summary
    doc.add_heading("Table of Compiled Code Modules", level=1)
    tbl_toc = doc.add_table(rows=len(FILES_MANIFEST) + 1, cols=4)
    tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["#", "Module Name", "File Path", "Language"]
    hdr_row = tbl_toc.rows[0]
    for j, h in enumerate(headers):
        c = hdr_row.cells[j]
        set_cell_background(c, "0F172A")
        set_cell_margins(c, 80, 80, 100, 100)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for i, item in enumerate(FILES_MANIFEST):
        row = tbl_toc.rows[i + 1]
        bg = "F8FAFC" if i % 2 == 0 else "FFFFFF"
        for j, text in enumerate([str(i + 1), item["section"].split("—")[0].strip(), item["path"], item["lang"]]):
            c = row.cells[j]
            set_cell_background(c, bg)
            set_cell_margins(c, 60, 60, 80, 80)
            p = c.paragraphs[0]
            r = p.add_run(text)
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_page_break()

    # Iterate and write all files in full
    for idx, item in enumerate(FILES_MANIFEST):
        fpath = item["path"]
        print(f"[DOCX] Processing ({idx+1}/{len(FILES_MANIFEST)}): {fpath}...")
        
        # Heading
        h = doc.add_heading(item["section"], level=1)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(2)

        # Meta line
        p_meta = doc.add_paragraph()
        p_meta.paragraph_format.space_after = Pt(6)
        r_path = p_meta.add_run(f"File Path: {fpath} | Language: {item['lang']}\n")
        r_path.font.name = "Calibri"
        r_path.font.size = Pt(9.5)
        r_path.font.bold = True
        r_path.font.color.rgb = RGBColor(14, 116, 144)

        r_desc = p_meta.add_run(item["description"])
        r_desc.font.name = "Calibri"
        r_desc.font.size = Pt(9)
        r_desc.font.italic = True
        r_desc.font.color.rgb = RGBColor(71, 85, 105)

        if not os.path.exists(fpath):
            p_missing = doc.add_paragraph()
            r_m = p_missing.add_run(f"[File not found at: {fpath}]")
            r_m.font.color.rgb = RGBColor(239, 68, 68)
            continue

        with open(fpath, "r", encoding="utf-8", errors="replace") as f:
            lines = f.readlines()

        # Format code block into shaded table
        tbl_code = doc.add_table(rows=1, cols=1)
        tbl_code.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl_code.rows[0].cells[0]
        cell.width = Inches(7.0)
        set_cell_background(cell, "0F172A") # Dark slate terminal theme
        set_cell_margins(cell, 120, 120, 140, 140)

        p_code = cell.paragraphs[0]
        p_code.paragraph_format.space_before = Pt(0)
        p_code.paragraph_format.space_after = Pt(0)
        p_code.paragraph_format.line_spacing = Pt(10.5)

        # Prepare numbered text
        formatted_code = []
        for line_num, line in enumerate(lines, 1):
            formatted_code.append(f"{line_num:4d} |  {line.rstrip()}")

        full_code_text = "\n".join(formatted_code)
        
        # Word paragraph runs limit ~ 65,000 chars, chunk if necessary
        chunk_size = 40000
        for i in range(0, len(full_code_text), chunk_size):
            chunk = full_code_text[i:i+chunk_size]
            r_code = p_code.add_run(chunk)
            r_code.font.name = "Consolas"
            r_code.font.size = Pt(8.0)
            r_code.font.color.rgb = RGBColor(226, 232, 240) # Slate-200

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

    output_path = "Source_Code_Compendium.docx"
    doc.save(output_path)
    print(f"[DOCX] Successfully saved: {output_path} ({os.path.getsize(output_path):,} bytes)")

# -------------------------------------------------------------
# Part 2: PDF Generation via ReportLab
# -------------------------------------------------------------
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image, Preformatted
)
from reportlab.pdfgen import canvas

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        
        # Running header (from page 2 onwards)
        if self._pageNumber > 1:
            self.drawString(54, letter[1] - 36, f"Project 32: {PROJECT_TITLE} — Complete Source Code")
            self.setStrokeColor(colors.HexColor("#CBD5E1"))
            self.setLineWidth(0.5)
            self.line(54, letter[1] - 42, letter[0] - 54, letter[1] - 42)

        # Running footer
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(letter[0] - 54, 30, page_str)
        self.drawString(54, 30, "HCLTech Career Shaper Programme | Learner: Saksham")
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(54, 42, letter[0] - 54, 42)
        
        self.restoreState()

def generate_pdf():
    print("[PDF] Initializing ReportLab PDF generator...")
    pdf_path = "Source_Code_Compendium.pdf"
    
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        leftMargin=40,
        rightMargin=40,
        topMargin=50,
        bottomMargin=50
    )

    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0F172A'),
        alignment=1, # Center
        spaceAfter=6
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#475569'),
        alignment=1,
        spaceAfter=14
    )

    h1_style = ParagraphStyle(
        'SectionH1',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#0F172A'),
        spaceBefore=14,
        spaceAfter=4,
        keepWithNext=True
    )

    desc_style = ParagraphStyle(
        'FileDesc',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#475569'),
        spaceAfter=8,
        keepWithNext=True
    )

    code_style = ParagraphStyle(
        'CodeStyle',
        fontName='Courier',
        fontSize=6.5,
        leading=8.2,
        textColor=colors.HexColor('#0F172A'),
        backColor=colors.HexColor('#F8FAFC'),
        borderColor=colors.HexColor('#CBD5E1'),
        borderWidth=0.5,
        borderPadding=4
    )

    elements = []

    # Title & Metadata
    elements.append(Paragraph("PROJECT SOURCE CODE COMPENDIUM", title_style))
    elements.append(Paragraph(f"{PROJECT_TITLE} (Project No. 32)<br/>{PROGRAMME}", subtitle_style))

    # Meta Table
    meta_data = [
        [Paragraph("<b>Student / Learner:</b>", styles['Normal']), Paragraph(f"{LEARNER} (CSE / AIML)", styles['Normal'])],
        [Paragraph("<b>Project Coordinator:</b>", styles['Normal']), Paragraph(COORDINATOR, styles['Normal'])],
        [Paragraph("<b>GitHub Repository:</b>", styles['Normal']), Paragraph(f"<a href='{GITHUB_REPO}'>{GITHUB_REPO}</a>", styles['Normal'])],
        [Paragraph("<b>Live Vercel App:</b>", styles['Normal']), Paragraph(f"<a href='{LIVE_VERCEL}'>{LIVE_VERCEL}</a>", styles['Normal'])],
    ]
    t_meta = Table(meta_data, colWidths=[130, 400])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('LEFTPADDING', (0,0), (-1,-1), 8),
        ('RIGHTPADDING', (0,0), (-1,-1), 8),
    ]))
    elements.append(t_meta)
    elements.append(Spacer(1, 14))

    # Architecture Image
    if os.path.exists("visualizations/00_system_architecture.png"):
        elements.append(Paragraph("System Architecture & Data Flow Diagram (Level 0)", h1_style))
        elements.append(Spacer(1, 4))
        elements.append(Image("visualizations/00_system_architecture.png", width=6.8*inch, height=3.6*inch))
        elements.append(Spacer(1, 14))

    # Table of Contents
    elements.append(Paragraph("Table of Compiled Code Modules", h1_style))
    toc_data = [["#", "Module Name", "File Path", "Language"]]
    for idx, item in enumerate(FILES_MANIFEST):
        toc_data.append([
            str(idx + 1),
            item["section"].split("—")[0].strip(),
            item["path"],
            item["lang"]
        ])
    t_toc = Table(toc_data, colWidths=[25, 250, 180, 75])
    t_toc.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0F172A')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 8),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.HexColor('#F8FAFC'), colors.white]),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    elements.append(t_toc)
    elements.append(PageBreak())

    # Add each source file
    for idx, item in enumerate(FILES_MANIFEST):
        fpath = item["path"]
        print(f"[PDF] Processing ({idx+1}/{len(FILES_MANIFEST)}): {fpath}...")
        
        elements.append(Paragraph(item["section"], h1_style))
        meta_sub = f"<b>File:</b> {fpath} &nbsp;|&nbsp; <b>Language:</b> {item['lang']}<br/><i>{item['description']}</i>"
        elements.append(Paragraph(meta_sub, desc_style))

        if not os.path.exists(fpath):
            elements.append(Paragraph(f"<font color='red'>[File not found: {fpath}]</font>", styles['Normal']))
            continue

        with open(fpath, "r", encoding="utf-8", errors="replace") as f:
            lines = f.readlines()

        code_lines = []
        for line_num, line in enumerate(lines, 1):
            code_lines.append(f"{line_num:4d} | {line.rstrip()}")

        # Add code in 45-line blocks so ReportLab splits across pages gracefully
        chunk_size = 45
        for c_idx in range(0, len(code_lines), chunk_size):
            chunk = code_lines[c_idx:c_idx+chunk_size]
            raw_chunk = "\n".join(chunk)
            pref = Preformatted(raw_chunk, code_style)
            elements.append(pref)
            elements.append(Spacer(1, 3))
        elements.append(Spacer(1, 10))

    # Build document
    doc.build(elements, canvasmaker=NumberedCanvas)
    print(f"[PDF] Successfully built: {pdf_path} ({os.path.getsize(pdf_path):,} bytes)")

if __name__ == "__main__":
    generate_docx()
    generate_pdf()
    print("\n[ALL DONE] Generated both Source_Code_Compendium.docx and Source_Code_Compendium.pdf successfully!")
